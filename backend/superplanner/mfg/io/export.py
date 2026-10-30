"""Bundle -> xlsx / zip-of-CSV bytes and blank import templates (M1-DESIGN 8.3, 8.4, review R4).

xlsx (openpyxl `write_only=True`): one sheet per template sheet (all columns, items with their `attr:<name>` columns)
plus a README sheet. Text values are explicit string cells (`data_type = "s"`), so a value starting with `=` is never
a formula and round-trips unchanged. Dates are date cells (`yyyy-mm-dd`), datetimes ISO text with offset.

CSV zip: one `<sheet>.csv` per sheet (UTF-8 with BOM, `,`, CRLF), fixed member timestamps (1980-01-01), so the bytes
are deterministic. Text values starting with `=`, `+`, `-`, `@`, tab, CR or `'` get a leading `'` that the reader
strips again.
"""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, datetime
from typing import Any

from ..bundle import ATTR_PREFIX, SHEETS, MfgBundle, SheetSpec
from ..units import UOMS
from .read import CSV_ESCAPE_TRIGGERS
from .spec import column_synonyms

ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)


def sheet_columns(spec: SheetSpec, rows: list[dict[str, Any]]) -> list[str]:
    cols = list(spec.fields)
    if spec.attr_columns:
        extra = sorted({k for r in rows for k in r if k.startswith(ATTR_PREFIX)})
        cols += extra
    return cols


def _text(v: Any) -> str:
    if isinstance(v, list):
        return ";".join(str(x) for x in v)
    return str(v)


def csv_value(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(v) if isinstance(v, float) else str(v)
    if isinstance(v, datetime):
        return v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = _text(v)
    if s and s[0] in CSV_ESCAPE_TRIGGERS:
        return "'" + s
    return s


def bundle_to_csv_zip(bundle: MfgBundle) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for spec in SHEETS:
            if spec.name not in bundle:          # scope=master / transactions exports (M2-DESIGN 11)
                continue
            rows = bundle.get(spec.name) or []
            cols = sheet_columns(spec, rows)
            out = io.StringIO(newline="")
            w = csv.writer(out, lineterminator="\r\n")
            w.writerow(cols)
            for r in rows:
                w.writerow([csv_value(r.get(c)) for c in cols])
            info = zipfile.ZipInfo(f"{spec.name}.csv", date_time=ZIP_EPOCH)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zf.writestr(info, ("﻿" + out.getvalue()).encode("utf-8"))
    return buf.getvalue()


def readme_rows() -> list[list[Any]]:
    rows: list[list[Any]] = [["sheet", "field", "type", "required on create", "key", "allowed values", "synonyms"]]
    for spec in SHEETS:
        for c in spec.columns:
            rows.append([spec.name, c.field, c.type, "yes" if c.required_on_create else "", "yes" if c.key else "",
                         ", ".join(c.enum_values) if c.enum_values else "", ", ".join(column_synonyms(c.field))])
        if spec.attr_columns:
            rows.append([spec.name, "attr:<name>", "str", "", "", "free attribute columns", ""])
    rows.append([])
    rows.append(["units", "code", "dimension", "factor", "", "aliases", ""])
    for u in UOMS:
        rows.append(["units", u.code, u.dimension, u.factor, "", ", ".join(u.aliases), ""])
    rows.append([])
    rows.append(["notes", "Lists are separated by ';' or ','. Booleans: true/false, yes/no, y/n, 1/0. Dates: "
                 "YYYY-MM-DD. Exception start/end: YYYY-MM-DD (local midnight; a date-only end includes that day), "
                 "YYYY-MM-DD HH:MM (local) or ISO with offset."])
    return rows


def bundle_to_xlsx(bundle: MfgBundle, *, readme: bool = True) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell

    wb = Workbook(write_only=True)

    def cell(ws: Any, v: Any) -> Any:
        if v is None:
            return None
        if isinstance(v, bool) or (isinstance(v, (int, float)) and not isinstance(v, bool)):
            return v
        if isinstance(v, datetime):
            c = WriteOnlyCell(ws, value=v.isoformat())
            c.data_type = "s"
            return c
        if isinstance(v, date):
            c = WriteOnlyCell(ws, value=v)
            c.number_format = "yyyy-mm-dd"
            return c
        c = WriteOnlyCell(ws, value=_text(v))
        c.data_type = "s"
        return c

    for spec in SHEETS:
        if spec.name not in bundle:
            continue
        ws = wb.create_sheet(spec.name)
        rows = bundle.get(spec.name) or []
        cols = sheet_columns(spec, rows)
        ws.append([cell(ws, c) for c in cols])
        for r in rows:
            ws.append([cell(ws, r.get(c)) for c in cols])
    if readme:
        ws = wb.create_sheet("README")
        for r in readme_rows():
            ws.append([cell(ws, v) for v in r])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_bundle() -> MfgBundle:
    """Blank templates: header rows plus one example row per sheet from the `small` sample."""
    from ..samples import build

    small = build("small")
    out: MfgBundle = {}
    for spec in SHEETS:
        rows = small.get(spec.name) or []
        out[spec.name] = rows[:1]
    return out
