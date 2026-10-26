"""15.3 #6: performance P2, P5, P7, P9 (M1-DESIGN 12) end to end through the HTTP API on the selected database.
PostgreSQL limits are x 1.5 except P7 (absolute). Measured numbers are printed with `-s` ([accm1-perf] lines).

P2  GET /items/{id}/explosion   Kavach < 1.0 s; kavach x2 (4,000+ lines) < 1.5 s
P5  GET /health                 Kavach < 2.0 s
P7  import commit of 5,000 rows (items + BOM lines) into an empty owner < 10 s end to end
P9  POST /samples/kavach/load   empty owner < 10 s
"""

from __future__ import annotations

import io
import time

import pytest
from fastapi.testclient import TestClient

from accm1_helpers import AS_OF, Api, make_app

pytestmark = pytest.mark.slow


def _limit(db_kind: str, s: float) -> float:
    return s * (1.5 if db_kind == "postgres" else 1.0)


def _t(fn):
    t = time.monotonic()
    out = fn()
    return out, time.monotonic() - t


def test_p9_p2_p5_kavach(api: Api, db_kind: str):
    out, p9 = _t(lambda: api.load("kavach"))
    assert out["health"]["ok"]
    top = api.item("KV-TPU-1000")
    api.get(f"/items/{top['id']}/explosion", on=AS_OF)                  # warm-up (first request imports code)
    ex, p2 = _t(lambda: api.get(f"/items/{top['id']}/explosion", on=AS_OF))
    assert ex["totals"]["rows"] >= 2300
    h, p5 = _t(lambda: api.get("/health", as_of=AS_OF))
    assert h["ok"]
    print(f"\n[accm1-perf] {db_kind} P9 kavach load {p9:.3f} s; P2 explosion {p2:.3f} s "
          f"({ex['totals']['rows']} rows); P5 health {p5:.3f} s")
    assert p9 < _limit(db_kind, 10.0), p9
    assert p2 < _limit(db_kind, 1.0), p2
    assert p5 < _limit(db_kind, 2.0), p5


def test_p2_kavach_x2(api: Api, db_kind: str):
    from superplanner.mfg import samples
    from superplanner.mfg.io import export

    data = export.bundle_to_xlsx(samples.scale("kavach", 2))           # test setup only: a 4,900-line data set
    rec, preview = api.wizard(data, "kavach-x2.xlsx", options={"replace_children": True, "release": True,
                                                               "default_effective_from": None})
    assert preview["error_count"] == 0, preview["row_errors"][:5]
    api.commit(rec)
    assert api.get("/summary")["counts"]["bom_lines"] >= 4000
    for pn in ("KV-TPU-1000", "KV2-TPU-1000"):
        top = api.item(pn)
        api.get(f"/items/{top['id']}/explosion", on=AS_OF)
        ex, dt = _t(lambda: api.get(f"/items/{top['id']}/explosion", on=AS_OF))
        print(f"\n[accm1-perf] {db_kind} P2 x2 {pn} {dt:.3f} s ({ex['totals']['rows']} rows)")
        assert dt < _limit(db_kind, 1.5), dt


def _xlsx_5000() -> bytes:
    """1,000 items (50 make) + 4,000 BOM lines written with openpyxl directly (an ERP-like file)."""
    import openpyxl

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Items")
    ws.append(["Part No", "Description", "Make/Buy", "UoM", "Lead time"])
    for i in range(50):
        ws.append([f"P7-MK-{i:03d}", f"Assembly {i}", "make", "EA", None])
    for i in range(950):
        ws.append([f"P7-BY-{i:04d}", f"Component {i}", "buy", "Nos", 14 + i % 30])
    ws = wb.create_sheet("BOM")
    ws.append(["Parent", "Find No", "Component", "Qty", "UoM"])
    for p in range(50):
        for k in range(80):
            ws.append([f"P7-MK-{p:03d}", str(10 * (k + 1)), f"P7-BY-{(p * 41 + k) % 950:04d}", 1 + k % 4, "EA"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_p7_import_commit_5000_rows_into_empty_owner(db_url_factory, db_kind: str):
    data = _xlsx_5000()
    with TestClient(make_app(db_url_factory("accm1_p7"))) as c:
        api = Api(c)
        rec, preview = api.wizard(data, "erp.xlsx")
        targets = {s["name"]: (s["suggested"] or {}).get("target") for s in rec["sheets"]}
        assert targets == {"Items": "items", "BOM": "bom_lines"}, targets
        assert preview["error_count"] == 0, preview["row_errors"][:5]
        done, p7 = _t(lambda: api.commit(rec))
        n = sum(v["create"] for v in done["result"]["counts"].values())
        print(f"\n[accm1-perf] {db_kind} P7 commit {p7:.3f} s ({n} rows created, server "
              f"{done['result']['duration_ms']} ms)")
        assert done["result"]["counts"]["items"]["create"] == 1000
        assert done["result"]["counts"]["bom_lines"]["create"] == 4000
        assert p7 < 10.0, p7
        # P8 (bonus): the same file again is all unchanged and fast
        rec2, preview2 = api.wizard(data, "erp.xlsx")
        assert all(v["create"] == v["update"] == v["delete"] == 0 for v in preview2["counts"].values())
        _, p8 = _t(lambda: api.commit(rec2))
        print(f"[accm1-perf] {db_kind} P8 re-import commit {p8:.3f} s")
        assert p8 < _limit(db_kind, 5.0)
