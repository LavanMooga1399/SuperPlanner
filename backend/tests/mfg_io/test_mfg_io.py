"""Pure import/export tests (M1-DESIGN 8, 15.1 `test_mfg_io.py`, review R4; PLAN-M1 B3-1). No database."""

from __future__ import annotations

import io
import itertools
import time
import zipfile
from datetime import date

import openpyxl
import pytest

from superplanner.mfg import limits as L
from superplanner.mfg import samples
from superplanner.mfg.bundle import bundle_to_dataset, dataset_to_bundle
from superplanner.mfg.dataset import MfgDataset
from superplanner.mfg.io import export, mapping, read
from superplanner.mfg.io.plan import ImportOptions, plan_import
from superplanner.mfg.model import PlantRecord

# ---------------------------------------------------------------- helpers


def _ids():
    c = itertools.count()
    return lambda prefix: f"{prefix}{next(c):032x}"


def empty() -> MfgDataset:
    return MfgDataset(plant=PlantRecord())


def xlsx(rows_by_sheet: dict[str, list[list[object]]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in rows_by_sheet.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def patch_zip(data: bytes, member: str, transform) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            body = src.read(info)
            if info.filename == member:
                body = transform(body)
            z.writestr(info.filename, body)
    return out.getvalue()


def map_suggested(raw: read.RawFile, **opts) -> tuple:
    m = {"sheets": {s["name"]: mapping.suggest_mapping(s) or {"target": None, "columns": {}} for s in raw.sheets},
         "options": opts}
    return mapping.apply_mapping(raw.sheets, m)


def plan_file(data: bytes, ds: MfgDataset | None = None, *, sheet: str | None = None, **opts):
    raw = read.read_upload(data, sheet=sheet)
    bundle, sources, errors = map_suggested(raw)
    p = plan_import(bundle, ds or empty(), ImportOptions(**opts), sources=sources, id_factory=_ids())
    return p, errors


# ---------------------------------------------------------------- format detection and safety


def test_defusedxml_is_active():
    assert openpyxl.DEFUSEDXML is True


def test_magic_detection_and_xls_refused():
    assert read.detect_format(b"PK\x03\x04rest") == "zip"
    assert read.detect_format(b"\xd0\xcf\x11\xe0" + b"\0" * 10) == "xls"
    assert read.detect_format(b"a,b\n1,2") == "csv"
    with pytest.raises(read.ImportFileInvalid, match=r"\.xls"):
        read.read_upload(b"\xd0\xcf\x11\xe0" + b"\0" * 100)
    with pytest.raises(read.ImportFileInvalid):
        read.read_upload(b"")
    with pytest.raises(read.ImportFileInvalid):
        read.read_upload(b"PK\x03\x04garbage-that-is-not-a-zip")


def test_zip_bomb_refused_before_decompression():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("items.csv", b"0" * (2 * 1024 * 1024))       # ratio ~ 1000
    with pytest.raises(read.ImportTooLarge) as e:
        read.read_upload(buf.getvalue())
    assert e.value.limit == "IMPORT_MAX_COMPRESSION_RATIO"
    # 1 KB compressed claiming 1 GB uncompressed: refused from the declared sizes alone
    data = bytearray(buf.getvalue())
    small = zipfile.ZipFile(io.BytesIO(bytes(data)))
    info = small.infolist()[0]
    fake = io.BytesIO()
    with zipfile.ZipFile(fake, "w", zipfile.ZIP_STORED) as z:
        zi = zipfile.ZipInfo("items.csv")
        z.writestr(zi, b"a,b\n")
    raw = bytearray(fake.getvalue())
    # patch the central-directory uncompressed size (offset 24 in the CD header) to 1 GiB
    cd = raw.rfind(b"PK\x01\x02")
    raw[cd + 24:cd + 28] = (1024 ** 3).to_bytes(4, "little")
    started = time.monotonic()
    with pytest.raises(read.ImportTooLarge):
        read.read_upload(bytes(raw))
    assert time.monotonic() - started < 1.0
    assert info.file_size > 0


def test_zip_member_count_and_unsafe_paths():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for i in range(L.IMPORT_MAX_ZIP_MEMBERS + 1):
            z.writestr(f"f{i}.csv", "a\n1\n")
    with pytest.raises(read.ImportTooLarge) as e:
        read.read_upload(buf.getvalue())
    assert e.value.limit == "IMPORT_MAX_ZIP_MEMBERS"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("../items.csv", "part_number\nX\n")
    with pytest.raises(read.ImportFileInvalid):
        read.read_upload(buf.getvalue())


def test_xlsx_far_row_and_wide_column_refused():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "items"
    ws.append(["part_number", "make_buy"])
    ws.cell(row=1_000_000, column=1, value="X")
    buf = io.BytesIO()
    wb.save(buf)
    started = time.monotonic()
    with pytest.raises(read.ImportTooLarge) as e:
        read.read_upload(buf.getvalue())
    assert e.value.limit == "IMPORT_MAX_SCANNED_ROWS" and time.monotonic() - started < 2.0
    wb = openpyxl.Workbook()
    wb.active.cell(row=1, column=200, value="far")
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(read.ImportTooLarge) as e:
        read.read_upload(buf.getvalue())
    assert e.value.limit == "IMPORT_MAX_COLUMNS"


def test_xlsx_inf_nan_control_chars_and_formulas():
    data = xlsx({"items": [["part_number", "make_buy", "lead_time_days", "description"],
                           ["P1", "buy", 5, "ok"], ["P2", "buy", 7, "badtext"], ["P3", "buy", 9, "plain"]]})

    def edit(body: bytes) -> bytes:
        s = body.decode()
        s = s.replace("<v>5</v>", "<v>1E999</v>", 1)
        s = s.replace("<v>9</v>", "<f>SUM(A1:A2)</f>", 1)
        return s.encode()

    assert read.stage_cell("bad\x01text") == ("badtext", "IMP_CONTROL_CHAR")
    assert read.stage_cell(float("inf")) == (None, "IMP_BAD_NUMBER")

    raw = read.read_upload(patch_zip(data, "xl/worksheets/sheet1.xml", edit))
    sh = raw.sheets[0]
    cells = [c for _n, row in sh["rows"] for c in row]
    assert all(c is None or isinstance(c, (str, int, float, bool)) for c in cells)
    assert not any(isinstance(c, str) and "\x01" in c for c in cells)
    codes = sorted(code for _r, _c, code in sh["issues"])
    assert codes == ["IMP_BAD_NUMBER", "IMP_FORMULA_NO_VALUE"]
    bundle, _src, errors = map_suggested(raw)
    assert sorted(e["code"] for e in errors) == ["IMP_BAD_NUMBER", "IMP_FORMULA_NO_VALUE"]
    assert {e["row"] for e in errors} == {2, 4}
    nan = patch_zip(data, "xl/worksheets/sheet1.xml", lambda b: b.replace(b"<v>5</v>", b"<v>NaN</v>", 1))
    with pytest.raises(read.ImportFileInvalid):
        read.read_upload(nan)
    nul = xlsx({"items": [["part_number"], ["A_x0000_B"]]})
    raw = read.read_upload(nul)
    assert "\x00" not in str(raw.sheets[0]["rows"])


def test_billion_laughs_refused():
    data = xlsx({"items": [["part_number"], ["X"]]})
    dtd = b'<!DOCTYPE lolz [<!ENTITY lol "lol"><!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">]>'

    def inject(body: bytes) -> bytes:
        return dtd + body.replace(b">X<", b">&lol2;<")

    evil = patch_zip(data, "xl/worksheets/sheet1.xml", inject)
    assert b"&lol2;" in zipfile.ZipFile(io.BytesIO(evil)).read("xl/worksheets/sheet1.xml")
    with pytest.raises(read.ImportFileInvalid):
        read.read_upload(evil)


def test_staged_size_and_cell_limit():
    with pytest.raises(read.ImportTooLarge) as e:
        read.read_upload(("part_number\n" + "x" * (L.IMPORT_MAX_CELL_CHARS + 1) + "\n").encode())
    assert e.value.limit == "IMPORT_MAX_CELL_CHARS"
    assert read.staged_size([{"a": "é"}]) == len('[{"a":"é"}]')


# ---------------------------------------------------------------- csv


def test_csv_parsing_delimiters_encodings_escapes():
    raw = read.read_upload("part_number;make_buy;description\nA;buy;'=cmd\n".encode("cp1252") + b"B;buy;caf\xe9\n",
                           sheet="items")
    assert raw.format == "csv" and raw.sheets[0]["name"] == "items"
    assert raw.warnings[0]["code"] == "IMP_ENCODING_FALLBACK"
    rows = [cells for _n, cells in raw.sheets[0]["rows"]]
    assert rows == [["A", "buy", "=cmd"], ["B", "buy", "café"]]
    raw = read.read_upload(b"\xef\xbb\xbfpart_number\tmake_buy\nA\tbuy\n\n\nB\tbuy\n")
    assert [n for n, _c in raw.sheets[0]["rows"]] == [2, 5]
    raw = read.read_upload(b"part_number,description\nA,x\x00y\n")
    assert raw.sheets[0]["issues"] == [[2, 1, "IMP_CONTROL_CHAR"]]


def test_cell_types_lists_booleans_aliases_and_numbers():
    data = xlsx({"items": [["Part No", "Make or buy", "UoM", "Lead time", "serialized", "MOQ", "desc"],
                           [12345, "Purchased", "pcs.", 14, "yes", 100.0, "a"],
                           ["B-1", "M", "MTR", None, "n", 0, "b"],
                           ["B-2", "buy", "furlong", "1,000", "maybe", "x", "c"]]})
    p, errors = plan_file(data)
    assert {(e["row"], e["code"]) for e in errors} == {(4, "IMP_BAD_UOM"), (4, "IMP_BAD_NUMBER"),
                                                       (4, "IMP_BAD_BOOLEAN")}
    created = {c.values["part_number"]: c.values for c in p.changes if c.table == "items"}
    assert created["12345"]["make_buy"] == "buy" and created["12345"]["uom"] == "EA"
    assert created["12345"]["serialized"] is True and created["12345"]["min_order_qty"] == 100
    assert created["B-1"]["make_buy"] == "make" and created["B-1"]["uom"] == "M"
    assert errors[0]["column"] in ("UoM", "Lead time", "serialized", "MOQ")


def test_mapping_suggestions_for_erp_headers():
    s = {"name": "BOM", "headers": ["Parent", "Find No", "Component", "Qty", "UoM", "Ref Des"], "rows": []}
    m = mapping.suggest_mapping(s)
    assert m == {"target": "bom_lines", "columns": {
        "Parent": "parent_part_number", "Find No": "find_number", "Component": "child_part_number",
        "Qty": "qty_per", "UoM": "uom", "Ref Des": "ref_designators"}}
    s = {"name": "Sheet1", "headers": ["Material", "Description", "Procurement type", "Lead time days"], "rows": []}
    m = mapping.suggest_mapping(s)
    assert m["target"] == "items" and m["columns"]["Material"] == "part_number"
    assert m["columns"]["Lead time days"] == "lead_time_days"
    assert mapping.suggest_mapping({"name": "README", "headers": ["sheet", "field", "type"], "rows": []}) is None
    staged = [{"name": "S", "headers": ["a", "b"], "rows": [], "issues": []}]
    with pytest.raises(mapping.MappingError) as e:
        mapping.validate_mapping(staged, {"sheets": {"S": {"target": "items", "columns": {
            "a": "part_number", "b": "part_number"}}}})
    assert e.value.errors[0]["type"] == "duplicate_field"
    with pytest.raises(mapping.MappingError):
        mapping.validate_mapping(staged, {"sheets": {"S": {"target": "nope", "columns": {}}}})
    with pytest.raises(mapping.MappingError):
        mapping.validate_mapping(staged, {"sheets": {"S": {"target": "items", "columns": {"zz": "part_number"}}}})


def test_excel_dates_and_exception_datetimes():
    from datetime import datetime

    data = xlsx({"plant": [["name", "utc_offset"], ["P", "+05:30"]],
                 "groups": [["code", "name"], ["G", "G"]],
                 "calendar_exceptions": [["scope", "target_code", "kind", "start", "end", "label"],
                                         ["plant", None, "holiday", datetime(2026, 11, 9), datetime(2026, 11, 9), "H"],
                                         ["group", "G", "maintenance", "2026-11-14 06:00", "2026-11-14 18:00", "M"]]})
    p, errors = plan_file(data)
    assert errors == [] and p.ok, p.row_errors
    exc = [c.values for c in p.changes if c.table == "calendar_exceptions"]
    assert exc[0]["start_at"].isoformat() == "2026-11-08T18:30:00+00:00"
    assert (exc[0]["end_at"] - exc[0]["start_at"]).days == 1                 # date-only end is inclusive
    assert exc[1]["start_at"].hour == 0 and exc[1]["start_at"].minute == 30


# ---------------------------------------------------------------- plan


def test_plan_diff_create_update_unchanged_delete_and_idempotency():
    b = samples.build("small")
    ds = bundle_to_dataset(b)
    p = plan_import(b, ds)
    assert p.ok and not p.has_writes()
    assert all(c["create"] == c["update"] == c["delete"] == 0 for c in p.counts.values())
    b2 = samples.build("small")
    b2["items"][0]["description"] = "changed"
    b2["bom_lines"] = b2["bom_lines"][1:]
    b2["bom_lines"].append(dict(b2["bom_lines"][0], find_number="999"))
    b2["boms"][0]["status"] = "draft"
    p = plan_import(b2, ds)
    assert p.counts["items"]["update"] == 1
    # the BOM is released: changing its lines is refused
    assert [e["code"] for e in p.row_errors] == ["IMP_RELEASED_READONLY"]
    # same change on a draft document is a diff
    b3 = samples.build("small")
    for r in b3["boms"] + b3["routings"]:
        r["status"] = "draft"
    ds3 = bundle_to_dataset(b3)
    b4 = samples.build("small")
    for r in b4["boms"] + b4["routings"]:
        r["status"] = "draft"
    b4["bom_lines"] = b4["bom_lines"][1:]
    b4["bom_lines"][0]["qty_per"] = 7.0
    p = plan_import(b4, ds3)
    assert p.ok and p.counts["bom_lines"]["delete"] == 1 and p.counts["bom_lines"]["update"] == 1
    assert p.bump["boms"] and any(s["action"] == "update" and s["fields"].get("qty_per") for s in p.changes_sample)
    p = plan_import(b4, ds3, ImportOptions(replace_children=False))
    assert p.counts["bom_lines"]["delete"] == 0


def test_plan_unknown_reference_duplicate_cycle_and_auto_find_numbers():
    csv = (b"parent_part_number,child_part_number,qty_per\nA,B,1\nA,C,2\nD,GHOST,1\nB,A,1\n")
    items = b"part_number,make_buy,lead_time_days\nA,make,\nB,make,\nC,buy,3\nC,buy,3\nD,make,\n"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("items.csv", items)
        z.writestr("bom_lines.csv", csv)
    p, errors = plan_file(buf.getvalue())
    codes = sorted((e["sheet"], e["code"]) for e in p.row_errors + errors)
    assert ("bom_lines", "IMP_UNKNOWN_REFERENCE") in codes and ("items", "IMP_DUPLICATE_KEY") in codes
    assert ("bom_lines", "IMP_BOM_CYCLE") in codes
    ghost = next(e for e in p.row_errors if e["code"] == "IMP_UNKNOWN_REFERENCE")
    assert ghost["row"] == 4 and ghost["column"] == "child_part_number" and "GHOST" in ghost["message"]
    lines = [c.values for c in p.changes if c.table == "bom_lines"]
    assert [ln["find_number"] for ln in lines if ln["qty_per"] in (1.0, 2.0)][:2] == ["10", "20"]


def test_plan_release_rules_on_create():
    b = samples.build("small")
    b["bom_lines"][0]["uom"] = "M"                           # count child with a length unit
    p = plan_import(b, empty(), ImportOptions(release=True))
    assert any(e["code"] == "IMP_RELEASE_BLOCKED" and "MFG_UOM_MISMATCH" in e["message"] for e in p.row_errors)
    b = samples.build("small")
    p = plan_import(b, empty(), ImportOptions(release=True))
    assert p.ok
    assert {c.values["status"] for c in p.changes if c.table in ("boms", "routings")} == {"released"}


def test_plan_shift_templates_assignments_operators():
    b = samples.build("small")
    ds = bundle_to_dataset(b)
    b["shifts"] = [r for r in b["shifts"] if r["template_code"] != "EVE"] + [
        {"template_code": "EVE", "shift_name": "Late", "days": ["mon"], "start": "18:00", "end": "23:00",
         "breaks": []}]
    p = plan_import(b, ds)
    assert p.counts["shifts"]["create"] == 1 and p.counts["shifts"]["delete"] == 1
    # assignments referencing the removed shift name -> row error
    asg = [r for r in b["shift_assignments"] if r["template_code"] == "EVE"]
    if asg:
        asg[0]["shift_names"] = ["Evening"]
        p = plan_import(b, ds)
        assert any(e["sheet"] == "shift_assignments" for e in p.row_errors)
    b = samples.build("small")
    b["shifts"].append({"template_code": "REG", "shift_name": "Overlap", "days": ["mon"], "start": "10:00",
                        "end": "11:00", "breaks": []})
    p = plan_import(b, ds)
    assert any(e["sheet"] == "shifts" for e in p.row_errors)


# ---------------------------------------------------------------- export


def test_xlsx_formula_text_round_trips_as_string_cell():
    b = samples.build("small")
    b["items"][0]["description"] = '=HYPERLINK("http://evil.example","x")'
    data = export.bundle_to_xlsx(b)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["items"]
    cell = ws.cell(row=2, column=3)
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    raw = read.read_upload(data)
    bundle, _s, errors = map_suggested(raw)
    assert errors == [] and bundle["items"][0]["description"] == b["items"][0]["description"]


def test_csv_quote_prefix_round_trips():
    b = samples.build("small")
    b["items"][0]["description"] = "=SUM(A1)"
    b["items"][1]["description"] = "'=already"
    b["items"][2]["description"] = "-5 dB"
    data = export.bundle_to_csv_zip(b)
    text = zipfile.ZipFile(io.BytesIO(data)).read("items.csv").decode("utf-8-sig")
    assert "'=SUM(A1)" in text
    bundle, _s, errors = map_suggested(read.read_upload(data))
    assert errors == []
    assert [r["description"] for r in bundle["items"][:3]] == ["=SUM(A1)", "'=already", "-5 dB"]


@pytest.mark.parametrize("sid", ["small", "medium", "kavach"])
def test_export_import_round_trip(sid: str):
    b = samples.build(sid)
    ds = bundle_to_dataset(b)
    for fmt in ("xlsx", "zip"):
        data = export.bundle_to_xlsx(b) if fmt == "xlsx" else export.bundle_to_csv_zip(b)
        p, errors = plan_file(data, ds)
        assert errors == [] and p.ok and not p.has_writes(), (fmt, p.row_errors[:3], p.changes[:3])
        p, errors = plan_file(data)
        assert errors == [] and p.ok, (fmt, p.row_errors[:3])
    assert export.bundle_to_csv_zip(b) == export.bundle_to_csv_zip(dataset_to_bundle(ds))


def test_templates_have_headers_and_readme():
    data = export.bundle_to_xlsx(export.template_bundle())
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    assert "README" in wb.sheetnames and "bom_lines" in wb.sheetnames
    z = zipfile.ZipFile(io.BytesIO(export.bundle_to_csv_zip(export.template_bundle())))
    assert {i.date_time for i in z.infolist()} == {(1980, 1, 1, 0, 0, 0)}
    assert date(2026, 1, 1)
